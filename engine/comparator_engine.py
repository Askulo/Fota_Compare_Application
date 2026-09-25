"""
FOTA Pre/Post Firmware Comparison Engine
-----------------------------------------
This is the original comparison + Excel-report logic, unchanged in behavior.
Only the Tkinter dialog-driven main() has been removed — everything here is
now called directly by the PySide6 UI (see ui/comparison_page.py).
"""

import os
import re
import hashlib
from datetime import datetime
from collections import defaultdict
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

APP_NAME = "FOTA Pre/Post Firmware Full Comparator"
COMPARATOR_VERSION = "3.0"

BLUE = "1F4E78"
DARK_BLUE = "17365D"
RED = "F4CCCC"
DARK_RED = "C00000"
YELLOW = "FFF2CC"
GREEN = "D9EAD3"
WHITE = "FFFFFF"


# ============================================================
# GENERAL HELPERS
# ============================================================

def clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\r", " ").replace("\n", " ").split())


def local_name(tag):
    if not tag:
        return ""
    return tag.split("}", 1)[-1]


def sha256_file(path):
    sha = hashlib.sha256()
    with open(path, "rb") as f:
        while True:
            chunk = f.read(1024 * 1024)
            if not chunk:
                break
            sha.update(chunk)
    return sha.hexdigest()


def safe_xml_parse(path):
    try:
        tree = ET.parse(path)
        return tree.getroot()
    except ET.ParseError as e:
        raise RuntimeError(f"XML parsing failed:\n\n{path}\n\n{e}")


# ============================================================
# XML FLATTEN
# ============================================================

def flatten_element(element, prefix, output):
    for attr_name, attr_value in sorted(element.attrib.items()):
        key = f"{prefix}.@{local_name(attr_name)}"
        output[key] = clean(attr_value)

    children = list(element)

    if not children:
        output[prefix] = clean(element.text)
        return

    direct_text = clean(element.text)
    if direct_text:
        output[f"{prefix}.__text"] = direct_text

    total_counts = defaultdict(int)
    for child in children:
        total_counts[local_name(child.tag)] += 1

    running_counts = defaultdict(int)
    for child in children:
        child_name = local_name(child.tag)
        running_counts[child_name] += 1
        count = running_counts[child_name]
        if total_counts[child_name] > 1:
            child_prefix = f"{prefix}.{child_name}[{count}]"
        else:
            child_prefix = f"{prefix}.{child_name}"
        flatten_element(child, child_prefix, output)


# ============================================================
# PARSE METER FILE
# ============================================================

def parse_meter_file(path):
    root = safe_xml_parse(path)
    objects = []
    children = list(root)

    for object_index, obj in enumerate(children, start=1):
        obj_class = local_name(obj.tag)
        record = {
            "ObjectIndex": object_index,
            "Class": obj_class,
            "LN": "",
            "Fields": {},
            "Raw": clean(ET.tostring(obj, encoding="unicode")),
        }

        for attr_name, attr_value in sorted(obj.attrib.items()):
            key = "@" + local_name(attr_name)
            record["Fields"][key] = clean(attr_value)

        for child in obj:
            flatten_element(child, local_name(child.tag), record["Fields"])

        ln_candidates = [
            "LN", "LogicalName", "Logical_Name", "logical_name",
            "LogicalName.Value", "Logical_Name.Value",
            "LogicalName.__text", "Logical_Name.__text",
        ]
        for candidate in ln_candidates:
            if candidate in record["Fields"]:
                value = clean(record["Fields"][candidate])
                if value:
                    record["LN"] = value
                    break

        if not record["LN"]:
            for field, value in record["Fields"].items():
                field_lower = field.lower()
                if "logical" in field_lower and "name" in field_lower and value:
                    record["LN"] = clean(value)
                    break
                if field_lower.endswith(".ln") and value:
                    record["LN"] = clean(value)
                    break

        if not record["LN"]:
            record["LN"] = f"__NO_LN_OBJECT_{object_index}"

        objects.append(record)

    return objects


# ============================================================
# OBJECT KEY / MAP
# ============================================================

def object_key(obj):
    return (clean(obj.get("Class")).lower(), clean(obj.get("LN")).lower())


def make_object_map(objects):
    base = defaultdict(list)
    for obj in objects:
        base[object_key(obj)].append(obj)

    result = {}
    for key, items in base.items():
        if len(items) == 1:
            result[key] = items[0]
        else:
            for index, item in enumerate(items, start=1):
                result[(key[0], key[1], f"DUPLICATE_{index}")] = item
    return result


def normalize_for_comparison(value):
    return clean(value)


# ============================================================
# CHANGE CLASSIFICATION
# ============================================================

def classify_change(field, old_value, new_value, old_obj, new_obj):
    field_lower = field.lower()
    obj = new_obj or old_obj
    class_name = clean(obj.get("Class")).lower()
    ln = clean(obj.get("LN")).lower()

    firmware_words = ["firmware", "softwareversion", "software_version",
                       "fwversion", "fw_version", "firmwareversion", "firmware_version"]
    if any(w in field_lower for w in firmware_words):
        return "FIRMWARE VERSION"
    if ln == "1.0.0.2.0.255":
        return "FIRMWARE VERSION"

    scaler_words = ["scaler", "scale", "unit", "scal_unit", "scalerunit", "scaler_unit"]
    if any(w in field_lower for w in scaler_words):
        return "SCALER / UNIT CHANGE"

    access_words = ["access", "accessmode", "access_mode", "attributeaccess",
                     "attribute_access", "methodaccess", "method_access",
                     "accessright", "access_right"]
    if any(w in field_lower for w in access_words):
        return "ACCESS / METHOD ACCESS CHANGE"

    security_words = ["security", "authentication", "auth", "password", "secret",
                       "key", "cipher", "encryption", "securitysuite", "security_suite",
                       "dedicatedkey", "dedicated_key", "globalkey", "global_key"]
    if any(w in field_lower for w in security_words) or "security" in class_name or "association" in class_name:
        return "SECURITY / ASSOCIATION CHANGE"

    communication_words = ["communication", "comm", "network", "apn", "ip", "port",
                            "server", "host", "address", "baud", "modem", "gprs",
                            "gsm", "tcp", "udp", "ppp", "ethernet", "wifi", "dns"]
    if any(w in field_lower for w in communication_words):
        return "COMMUNICATION CONFIGURATION CHANGE"

    communication_classes = ["modem", "tcp", "udp", "ppp", "ethernet", "gprs", "gsm", "communication"]
    if any(w in class_name for w in communication_classes):
        return "COMMUNICATION CONFIGURATION CHANGE"

    parameter_words = ["config", "configuration", "parameter", "param", "setting",
                        "threshold", "limit", "enable", "disable", "mode", "interval",
                        "period", "timeout", "retry", "retries", "cycle", "frequency",
                        "delay", "offset"]
    if any(w in field_lower for w in parameter_words):
        return "CONFIGURATION / PARAMETER CHANGE"

    fota_words = ["fota", "image", "transfer", "block", "activate", "verify",
                  "imageactivate", "image_transfer"]
    if any(w in field_lower for w in fota_words):
        return "FOTA / IMAGE TRANSFER CHANGE"
    if "imagetransfer" in class_name or "image" in class_name:
        return "FOTA / IMAGE TRANSFER CHANGE"

    profile_words = ["profile", "capture", "captureobject", "capture_object",
                      "period", "sortmethod", "sort_method"]
    if any(w in field_lower for w in profile_words):
        return "PROFILE / CAPTURE CONFIGURATION"

    event_words = ["event", "action", "actionset", "action_set", "eventcode", "event_code"]
    if any(w in field_lower for w in event_words):
        return "EVENT / ACTION CONFIGURATION"

    time_words = ["clock", "time", "timezone", "time_zone", "deviation", "dst", "daylight"]
    if any(w in field_lower for w in time_words):
        return "DATE / TIME / TIMEZONE CHANGE"

    counter_words = ["counter", "invocation", "entriesinuse", "entries_in_use"]
    if any(w in field_lower for w in counter_words):
        return "COUNTER / RUNTIME STATE CHANGE"

    register_classes = ["register", "extendedregister", "demandregister", "registeractivation"]
    if any(x in class_name for x in register_classes):
        return "REGISTER / RUNTIME VALUE CHANGE"

    manufacturer_words = ["manufacturer", "vendor", "private", "custom", "specific", "proprietary"]
    if any(w in field_lower for w in manufacturer_words):
        return "MANUFACTURER-SPECIFIC CHANGE"

    return "OTHER VALUE / ATTRIBUTE CHANGE"


def severity(category):
    high_categories = {
        "FIRMWARE VERSION", "SCALER / UNIT CHANGE", "ACCESS / METHOD ACCESS CHANGE",
        "CONFIGURATION / PARAMETER CHANGE", "SECURITY / ASSOCIATION CHANGE",
        "COMMUNICATION CONFIGURATION CHANGE", "PROFILE / CAPTURE CONFIGURATION",
        "EVENT / ACTION CONFIGURATION", "FOTA / IMAGE TRANSFER CHANGE",
        "OBJECT STRUCTURE CHANGE",
    }
    medium_categories = {
        "MANUFACTURER-SPECIFIC CHANGE", "DATE / TIME / TIMEZONE CHANGE",
        "COUNTER / RUNTIME STATE CHANGE", "REGISTER / RUNTIME VALUE CHANGE",
    }
    if category in high_categories:
        return "HIGH - REVIEW / APPROVAL"
    if category in medium_categories:
        return "MEDIUM - REVIEW"
    return "LOW - INFORMATIONAL"


# ============================================================
# FULL COMPARISON
# ============================================================

def compare_files(old_objects, new_objects):
    old_map = make_object_map(old_objects)
    new_map = make_object_map(new_objects)
    all_keys = sorted(set(old_map.keys()) | set(new_map.keys()), key=str)

    changes = []
    unchanged_objects = 0
    common_objects = 0
    added_objects = 0
    removed_objects = 0
    changed_objects = 0

    for key in all_keys:
        old_obj = old_map.get(key)
        new_obj = new_map.get(key)

        if old_obj is None:
            added_objects += 1
            changes.append({
                "ChangeType": "OBJECT ADDED", "Category": "OBJECT STRUCTURE CHANGE",
                "Severity": "HIGH - REVIEW / APPROVAL", "Class": new_obj["Class"],
                "LN": new_obj["LN"], "Field": "[OBJECT]", "OldValue": "",
                "NewValue": "[OBJECT PRESENT ONLY IN CANDIDATE]", "ObjectKey": str(key),
            })
            continue

        if new_obj is None:
            removed_objects += 1
            changes.append({
                "ChangeType": "OBJECT REMOVED", "Category": "OBJECT STRUCTURE CHANGE",
                "Severity": "HIGH - REVIEW / APPROVAL", "Class": old_obj["Class"],
                "LN": old_obj["LN"], "Field": "[OBJECT]",
                "OldValue": "[OBJECT PRESENT ONLY IN BASELINE]", "NewValue": "",
                "ObjectKey": str(key),
            })
            continue

        common_objects += 1
        object_changed = False
        all_fields = sorted(set(old_obj["Fields"].keys()) | set(new_obj["Fields"].keys()))

        for field in all_fields:
            old_value = normalize_for_comparison(old_obj["Fields"].get(field, ""))
            new_value = normalize_for_comparison(new_obj["Fields"].get(field, ""))
            if old_value == new_value:
                continue
            object_changed = True
            category = classify_change(field, old_value, new_value, old_obj, new_obj)
            changes.append({
                "ChangeType": "VALUE / ATTRIBUTE CHANGED", "Category": category,
                "Severity": severity(category), "Class": new_obj["Class"],
                "LN": new_obj["LN"], "Field": field, "OldValue": old_value,
                "NewValue": new_value, "ObjectKey": str(key),
            })

        if object_changed:
            changed_objects += 1
        else:
            unchanged_objects += 1

    return {
        "changes": changes, "common_objects": common_objects,
        "added_objects": added_objects, "removed_objects": removed_objects,
        "changed_objects": changed_objects, "unchanged_objects": unchanged_objects,
        "old_map": old_map, "new_map": new_map,
    }


# ============================================================
# FIND FIRMWARE / METER NUMBER
# ============================================================

def find_firmware(objects):
    candidates = []
    for obj in objects:
        for field, value in obj["Fields"].items():
            field_lower = field.lower()
            if ("firmware" in field_lower or "softwareversion" in field_lower
                    or "firmwareversion" in field_lower or "fwversion" in field_lower):
                if value:
                    candidates.append(clean(value))
        if clean(obj["LN"]) == "1.0.0.2.0.255":
            for value in obj["Fields"].values():
                if value:
                    candidates.append(clean(value))

    result, seen = [], set()
    for value in candidates:
        if value and value not in seen:
            seen.add(value)
            result.append(value)

    return " | ".join(result) if result else "NOT FOUND"


def find_meter_number(old_path, new_path, old_objects, new_objects):
    filenames = [os.path.basename(old_path), os.path.basename(new_path)]
    patterns = [r"\b[A-Z]{2,10}\d{5,20}\b", r"\b\d{8,20}\b"]

    for filename in filenames:
        for pattern in patterns:
            match = re.search(pattern, filename, re.IGNORECASE)
            if match:
                return match.group(0)

    for objects in [old_objects, new_objects]:
        for obj in objects:
            for field, value in obj["Fields"].items():
                field_lower = field.lower()
                if any(x in field_lower for x in
                       ["meterno", "meter_no", "meterid", "meter_id",
                        "serial", "serialnumber", "serial_number"]):
                    value = clean(value)
                    if value:
                        return value

    return "NOT IDENTIFIED"


# ============================================================
# REPORT STYLE HELPERS
# ============================================================

def apply_header_style(sheet, row, columns):
    for col in range(1, columns + 1):
        cell = sheet.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="808080"))


def autosize(sheet, maximum=65):
    for column in sheet.columns:
        column_letter = get_column_letter(column[0].column)
        values = [str(cell.value if cell.value is not None else "") for cell in column]
        width = max((len(v) for v in values), default=10) + 2
        sheet.column_dimensions[column_letter].width = min(width, maximum)


def style_all_cells(sheet):
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


# ============================================================
# WRITE EXCEL REPORT  (unchanged workbook structure/content)
# ============================================================

def write_excel_report(output_path, old_path, new_path, old_objects, new_objects,
                        comparison, old_hash, new_hash):
    changes = comparison["changes"]
    common_objects = comparison["common_objects"]
    added_objects = comparison["added_objects"]
    removed_objects = comparison["removed_objects"]
    changed_objects = comparison["changed_objects"]
    unchanged_objects = comparison["unchanged_objects"]

    meter_number = find_meter_number(old_path, new_path, old_objects, new_objects)
    old_firmware = find_firmware(old_objects)
    new_firmware = find_firmware(new_objects)

    category_counts = defaultdict(int)
    severity_counts = defaultdict(int)
    for change in changes:
        category_counts[change["Category"]] += 1
        severity_counts[change["Severity"]] += 1

    scaler_changes = [x for x in changes if x["Category"] == "SCALER / UNIT CHANGE"]
    parameter_changes = [x for x in changes if x["Category"] in (
        "CONFIGURATION / PARAMETER CHANGE", "COMMUNICATION CONFIGURATION CHANGE",
        "SECURITY / ASSOCIATION CHANGE", "PROFILE / CAPTURE CONFIGURATION",
        "EVENT / ACTION CONFIGURATION")]
    access_changes = [x for x in changes if x["Category"] == "ACCESS / METHOD ACCESS CHANGE"]
    fota_changes = [x for x in changes if x["Category"] in (
        "FIRMWARE VERSION", "FOTA / IMAGE TRANSFER CHANGE")]
    runtime_changes = [x for x in changes if x["Category"] in (
        "REGISTER / RUNTIME VALUE CHANGE", "COUNTER / RUNTIME STATE CHANGE",
        "DATE / TIME / TIMEZONE CHANGE")]

    wb = Workbook()

    # ---- 01 Executive Summary ----
    ws = wb.active
    ws.title = "01 Executive Summary"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "FOTA PRE / POST FIRMWARE FULL COMPARISON REPORT"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_BLUE)
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 32

    summary = [
        ("Meter Number", meter_number),
        ("Baseline File", os.path.basename(old_path)),
        ("Candidate File", os.path.basename(new_path)),
        ("Baseline Firmware", old_firmware),
        ("Candidate Firmware", new_firmware),
        ("Baseline SHA-256", old_hash),
        ("Candidate SHA-256", new_hash),
        ("Comparison Timestamp", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Comparator Version", COMPARATOR_VERSION),
    ]
    for r, (k, v) in enumerate(summary, start=3):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v).alignment = Alignment(wrap_text=True)

    start = 15
    ws.cell(start, 1, "OBJECT COMPARISON STATISTICS").font = Font(bold=True, size=13)
    object_stats = [
        ("Baseline Objects", len(old_objects)), ("Candidate Objects", len(new_objects)),
        ("Common Objects", common_objects), ("Unchanged Common Objects", unchanged_objects),
        ("Changed Objects", changed_objects), ("Objects Added", added_objects),
        ("Objects Removed", removed_objects),
    ]
    for r, (k, v) in enumerate(object_stats, start + 1):
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)

    change_start = start
    ws.cell(change_start, 4, "CHANGE STATISTICS").font = Font(bold=True, size=13)
    change_stats = [
        ("Total Detected Changes", len(changes)),
        ("HIGH Changes", severity_counts["HIGH - REVIEW / APPROVAL"]),
        ("MEDIUM Changes", severity_counts["MEDIUM - REVIEW"]),
        ("LOW Changes", severity_counts["LOW - INFORMATIONAL"]),
        ("Scaler / Unit Changes", len(scaler_changes)),
        ("Parameter / Configuration Changes", len(parameter_changes)),
        ("Access Changes", len(access_changes)),
        ("Firmware / FOTA Changes", len(fota_changes)),
        ("Runtime / Counter / Time Changes", len(runtime_changes)),
    ]
    for r, (k, v) in enumerate(change_stats, change_start + 1):
        ws.cell(r, 4, k).font = Font(bold=True)
        ws.cell(r, 5, v)
        if "HIGH" in k:
            ws.cell(r, 5).fill = PatternFill("solid", fgColor=RED)
        elif "MEDIUM" in k:
            ws.cell(r, 5).fill = PatternFill("solid", fgColor=YELLOW)

    final_row = 29
    ws.cell(final_row, 1, "AUTOMATED COMPARISON STATUS").font = Font(bold=True, size=13)
    if len(changes) == 0:
        status, status_fill = "NO DIFFERENCE DETECTED IN COMPARED EXPORT DATA", GREEN
    elif not (scaler_changes or parameter_changes or access_changes or fota_changes):
        status = ("DIFFERENCES DETECTED - NO SCALER / PARAMETER / ACCESS / FOTA "
                   "CHANGE DETECTED BY THIS COMPARISON")
        status_fill = YELLOW
    else:
        status = "CHANGES DETECTED - ENGINEERING REVIEW REQUIRED BEFORE FOTA APPROVAL"
        status_fill = RED

    ws.cell(final_row + 1, 1, status)
    ws.cell(final_row + 1, 1).fill = PatternFill("solid", fgColor=status_fill)
    ws.cell(final_row + 1, 1).font = Font(bold=True)
    ws.merge_cells(start_row=final_row + 1, start_column=1, end_row=final_row + 2, end_column=6)
    ws.cell(final_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="center")

    note_row = 33
    ws.cell(note_row, 1, "IMPORTANT NOTE").font = Font(bold=True, color=DARK_RED)
    ws.cell(note_row + 1, 1, (
        "This report records differences observed in the supplied meter export files. "
        "A runtime value difference is not by itself proof of a firmware-code change. "
        "Firmware release notes, test evidence, package/hash verification and formal "
        "engineering approval must be reviewed before production FOTA."))
    ws.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 3, end_column=6)
    ws.cell(note_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    autosize(ws)

    # ---- 02 Full Change Register ----
    ws = wb.create_sheet("02 FULL Change Register")
    ws.sheet_view.showGridLines = False
    headers = ["Change ID", "Change Type", "Severity", "Category", "Class",
               "Logical Name (LN)", "Field / Attribute", "Baseline Value",
               "Candidate Value", "Object Key", "Engineering Review",
               "Approval Status", "Remarks"]
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header)
    apply_header_style(ws, 1, len(headers))

    for row_number, change in enumerate(changes, start=2):
        row = [row_number - 1, change["ChangeType"], change["Severity"], change["Category"],
               change["Class"], change["LN"], change["Field"], change["OldValue"],
               change["NewValue"], change["ObjectKey"], "", "", ""]
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row_number, col, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if change["Severity"].startswith("HIGH"):
            for col in range(1, len(headers) + 1):
                ws.cell(row_number, col).fill = PatternFill("solid", fgColor=RED)
        elif change["Severity"].startswith("MEDIUM"):
            for col in range(1, len(headers) + 1):
                ws.cell(row_number, col).fill = PatternFill("solid", fgColor=YELLOW)

    ws.freeze_panes = "A2"
    if changes:
        ws.auto_filter.ref = f"A1:M{len(changes) + 1}"
    autosize(ws)

    # ---- Category-specific audit sheets ----
    audit_headers = ["Change ID", "Class", "Logical Name (LN)", "Field",
                      "Baseline", "Candidate", "Status", "Engineering Remarks"]

    def write_audit_sheet(title, rows, empty_message, status_label):
        sheet = wb.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        for col, header in enumerate(audit_headers, start=1):
            sheet.cell(1, col, header)
        apply_header_style(sheet, 1, len(audit_headers))
        if not rows:
            sheet.cell(2, 1, empty_message)
            sheet.cell(2, 1).fill = PatternFill("solid", fgColor=GREEN)
            sheet.merge_cells("A2:H2")
        else:
            for row_number, change in enumerate(rows, start=2):
                row = [row_number - 1, change["Class"], change["LN"], change["Field"],
                       change["OldValue"], change["NewValue"], status_label, ""]
                for col, value in enumerate(row, start=1):
                    sheet.cell(row_number, col, value)
                    if status_label != "OBSERVED - REVIEW IF REQUIRED":
                        sheet.cell(row_number, col).fill = PatternFill("solid", fgColor=RED)
        autosize(sheet)

    write_audit_sheet("03 Scaler Unit Audit", scaler_changes,
                       "NO SCALER / UNIT CHANGE DETECTED", "CHANGED - REVIEW REQUIRED")
    write_audit_sheet("04 Parameter Config Audit", parameter_changes,
                       "NO CONFIGURATION / PARAMETER CHANGE DETECTED", "CHANGED - REVIEW REQUIRED")

    access_security_changes = [x for x in changes if x["Category"] in (
        "ACCESS / METHOD ACCESS CHANGE", "SECURITY / ASSOCIATION CHANGE")]
    write_audit_sheet("05 Access Security Audit", access_security_changes,
                       "NO ACCESS / SECURITY CHANGE DETECTED", "CHANGED - REVIEW REQUIRED")
    write_audit_sheet("06 Firmware FOTA Audit", fota_changes,
                       "NO FIRMWARE / FOTA CHANGE DETECTED", "REVIEW REQUIRED")
    write_audit_sheet("07 Runtime Changes", runtime_changes,
                       "NO RUNTIME / COUNTER / TIME CHANGE DETECTED", "OBSERVED - REVIEW IF REQUIRED")

    # ---- 08 Object Inventory ----
    ws = wb.create_sheet("08 Object Inventory")
    ws.sheet_view.showGridLines = False
    inventory_headers = ["Class", "Logical Name (LN)", "Baseline Present", "Candidate Present",
                          "Object Status", "Baseline Fields", "Candidate Fields",
                          "Field Difference Count"]
    for col, header in enumerate(inventory_headers, start=1):
        ws.cell(1, col, header)
    apply_header_style(ws, 1, len(inventory_headers))

    old_map = comparison["old_map"]
    new_map = comparison["new_map"]
    all_keys = sorted(set(old_map.keys()) | set(new_map.keys()), key=str)

    for row_number, key in enumerate(all_keys, start=2):
        old_obj = old_map.get(key)
        new_obj = new_map.get(key)

        if old_obj and new_obj:
            fields = set(old_obj["Fields"].keys()) | set(new_obj["Fields"].keys())
            difference_count = sum(
                1 for f in fields
                if normalize_for_comparison(old_obj["Fields"].get(f, "")) !=
                normalize_for_comparison(new_obj["Fields"].get(f, "")))
            status = "CHANGED" if difference_count else "UNCHANGED"
        elif old_obj:
            status, difference_count = "REMOVED", len(old_obj["Fields"])
        else:
            status, difference_count = "ADDED", len(new_obj["Fields"])

        obj = new_obj or old_obj
        row = [obj["Class"], obj["LN"], "YES" if old_obj else "NO",
               "YES" if new_obj else "NO", status,
               len(old_obj["Fields"]) if old_obj else 0,
               len(new_obj["Fields"]) if new_obj else 0, difference_count]
        for col, value in enumerate(row, start=1):
            ws.cell(row_number, col, value)

        status_cell = ws.cell(row_number, 5)
        if status == "UNCHANGED":
            status_cell.fill = PatternFill("solid", fgColor=GREEN)
        elif status == "CHANGED":
            status_cell.fill = PatternFill("solid", fgColor=YELLOW)
        else:
            status_cell.fill = PatternFill("solid", fgColor=RED)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    # ---- 09 Approval Checklist ----
    ws = wb.create_sheet("09 Approval Checklist")
    ws.sheet_view.showGridLines = False
    checklist_headers = ["Checkpoint", "Automated Status", "Required Evidence / Action",
                          "Owner", "Approval", "Date", "Remarks"]
    for col, header in enumerate(checklist_headers, start=1):
        ws.cell(1, col, header)
    apply_header_style(ws, 1, len(checklist_headers))

    checklist = [
        ("Full object comparison completed", "PASS", "All detected objects and fields compared."),
        ("Scaler / Unit changes reviewed", "PASS" if not scaler_changes else "PENDING",
         "Review 03 Scaler Unit Audit."),
        ("Parameter / configuration changes reviewed", "PASS" if not parameter_changes else "PENDING",
         "Review 04 Parameter Config Audit."),
        ("Access changes reviewed", "PASS" if not access_changes else "PENDING",
         "Review access and method permissions."),
        ("Security changes reviewed",
         "PENDING" if any(x["Category"] == "SECURITY / ASSOCIATION CHANGE" for x in changes) else "PASS",
         "Verify security/association configuration."),
        ("Firmware version reviewed", "PENDING", "Verify approved vendor release notes."),
        ("FOTA package/hash verified", "PENDING", "Verify approved image/package SHA-256."),
        ("Profile/event changes reviewed", "PENDING", "Validate profile and event configuration."),
        ("Pilot meter validation", "PENDING", "Record pilot meter testing evidence."),
        ("HES validation", "PENDING", "Validate readings, profiles and events."),
        ("MDM validation", "PENDING", "Validate downstream data availability."),
        ("Rollback plan verified", "PENDING", "Record rollback firmware and procedure."),
        ("Final FOTA approval", "PENDING", "Formal approval before production FOTA."),
    ]
    for row_number, (checkpoint, status, action) in enumerate(checklist, start=2):
        row = [checkpoint, status, action, "", "", "", ""]
        for col, value in enumerate(row, start=1):
            ws.cell(row_number, col, value)
        ws.cell(row_number, 2).fill = PatternFill(
            "solid", fgColor=GREEN if status == "PASS" else YELLOW)
    autosize(ws)

    # ---- 10 Audit Trail ----
    ws = wb.create_sheet("10 Audit Trail")
    ws.sheet_view.showGridLines = False
    audit_rows = [
        ("Report Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Comparator", APP_NAME), ("Comparator Version", COMPARATOR_VERSION),
        ("Baseline File", os.path.abspath(old_path)), ("Candidate File", os.path.abspath(new_path)),
        ("Baseline SHA-256", old_hash), ("Candidate SHA-256", new_hash),
        ("Baseline Object Count", len(old_objects)), ("Candidate Object Count", len(new_objects)),
        ("Common Objects", common_objects), ("Changed Objects", changed_objects),
        ("Unchanged Objects", unchanged_objects), ("Added Objects", added_objects),
        ("Removed Objects", removed_objects), ("Total Field/Attribute Changes", len(changes)),
    ]
    ws.cell(1, 1, "Audit Item")
    ws.cell(1, 2, "Value")
    apply_header_style(ws, 1, 2)
    for row_number, (k, v) in enumerate(audit_rows, start=2):
        ws.cell(row_number, 1, k).font = Font(bold=True)
        ws.cell(row_number, 2, v)
    autosize(ws)

    # ---- 11 Change Category Summary ----
    ws = wb.create_sheet("11 Change Category Summary")
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Change Category")
    ws.cell(1, 2, "Count")
    apply_header_style(ws, 1, 2)
    category_row = 2
    for category, count in sorted(category_counts.items()):
        ws.cell(category_row, 1, category)
        ws.cell(category_row, 2, count)
        if any(w in category for w in
               ("SCALER", "CONFIGURATION", "SECURITY", "ACCESS", "FIRMWARE", "FOTA")):
            ws.cell(category_row, 2).fill = PatternFill("solid", fgColor=RED)
        category_row += 1
    autosize(ws)

    # ---- 12 Field Snapshot ----
    ws = wb.create_sheet("12 Field Snapshot")
    ws.sheet_view.showGridLines = False
    snapshot_headers = ["Class", "LN", "Field", "Baseline", "Candidate", "Status"]
    for col, header in enumerate(snapshot_headers, start=1):
        ws.cell(1, col, header)
    apply_header_style(ws, 1, len(snapshot_headers))

    row_number = 2
    for key in all_keys:
        old_obj = old_map.get(key)
        new_obj = new_map.get(key)
        if not old_obj or not new_obj:
            continue
        fields = sorted(set(old_obj["Fields"].keys()) | set(new_obj["Fields"].keys()))
        for field in fields:
            old_value = normalize_for_comparison(old_obj["Fields"].get(field, ""))
            new_value = normalize_for_comparison(new_obj["Fields"].get(field, ""))
            status = "UNCHANGED" if old_value == new_value else "CHANGED"
            row = [new_obj["Class"], new_obj["LN"], field, old_value, new_value, status]
            for col, value in enumerate(row, start=1):
                ws.cell(row_number, col, value)
            ws.cell(row_number, 6).fill = PatternFill(
                "solid", fgColor=YELLOW if status == "CHANGED" else GREEN)
            row_number += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        style_all_cells(sheet)

    wb.save(output_path)


# ============================================================
# HIGH-LEVEL ENTRY POINT — used by the UI
# ============================================================

def run_comparison(old_file, new_file, output_dir=None):
    """
    Runs the full pipeline (hash -> parse -> compare -> write report) and
    returns a dict with everything the UI needs to render the results screen.
    """
    if os.path.abspath(old_file) == os.path.abspath(new_file):
        raise RuntimeError("Baseline and Candidate files cannot be the same.")

    old_hash = sha256_file(old_file)
    new_hash = sha256_file(new_file)

    old_objects = parse_meter_file(old_file)
    new_objects = parse_meter_file(new_file)

    if not old_objects:
        raise RuntimeError("No XML objects found in baseline file.")
    if not new_objects:
        raise RuntimeError("No XML objects found in candidate file.")

    comparison = compare_files(old_objects, new_objects)

    meter_number = find_meter_number(old_file, new_file, old_objects, new_objects)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_meter = re.sub(r"[^A-Za-z0-9_.-]", "_", meter_number)
    output_name = f"FOTA_Comparison_{safe_meter}_{timestamp}.xlsx"

    if output_dir is None:
        output_dir = os.path.dirname(os.path.abspath(new_file))
    output_path = os.path.join(output_dir, output_name)

    write_excel_report(output_path, old_file, new_file, old_objects, new_objects,
                        comparison, old_hash, new_hash)

    changes = comparison["changes"]
    category_counts = defaultdict(int)
    for c in changes:
        category_counts[c["Category"]] += 1

    return {
        "old_file": old_file, "new_file": new_file,
        "old_hash": old_hash, "new_hash": new_hash,
        "old_objects": old_objects, "new_objects": new_objects,
        "meter_number": meter_number,
        "old_firmware": find_firmware(old_objects),
        "new_firmware": find_firmware(new_objects),
        "comparison": comparison,
        "changes": changes,
        "category_counts": dict(category_counts),
        "high": sum(1 for x in changes if x["Severity"].startswith("HIGH")),
        "medium": sum(1 for x in changes if x["Severity"].startswith("MEDIUM")),
        "low": sum(1 for x in changes if x["Severity"].startswith("LOW")),
        "output_path": output_path,
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
